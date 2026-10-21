"""
Automated Tests for Excel Import Validation, Template Generation, and Formula Injection Defense.
"""

import io
from datetime import date
import openpyxl
from django.test import TestCase
from django.contrib.auth.models import User
from attendance.models import (
    StudentProfile, ClassSection, Department, AcademicYear, Semester, Subject
)
from attendance.services.excel_service import (
    sanitize_cell_value, generate_attendance_import_template,
    parse_and_validate_attendance_excel
)


class ExcelImportExportTestCase(TestCase):
    def setUp(self):
        self.year = AcademicYear.objects.create(name="2026-27", start_date=date(2026, 1, 1), end_date=date(2026, 12, 31))
        self.sem = Semester.objects.create(academic_year=self.year, number=1, name="Sem 1")
        self.dept = Department.objects.create(code="AIDS", name="AI & Data Science")
        self.sec = ClassSection.objects.create(department=self.dept, semester=self.sem, name="AIDS-A")
        self.subject = Subject.objects.create(code="AI101", name="Intro to AI", department=self.dept, semester=self.sem)

        self.st_user = User.objects.create_user(username="test_st", password="password")
        self.student = StudentProfile.objects.create(user=self.st_user, roll_number="ST001", registration_number="REG001", class_section=self.sec)

    def test_formula_injection_defense(self):
        """Sanitizes spreadsheet formula injection characters (=, +, -, @)."""
        self.assertEqual(sanitize_cell_value("=1+1"), "'=1+1")
        self.assertEqual(sanitize_cell_value("+cmd|' /C calc'!A0"), "'+cmd|' /C calc'!A0")
        self.assertEqual(sanitize_cell_value("-5"), "'-5")
        self.assertEqual(sanitize_cell_value("@SUM(A1:A10)"), "'@SUM(A1:A10)")
        self.assertEqual(sanitize_cell_value("Normal Text"), "Normal Text")

    def test_template_generation_returns_valid_workbook(self):
        """Template generation produces valid openpyxl workbook stream with expected headers."""
        stream = generate_attendance_import_template(self.sec, self.subject)
        wb = openpyxl.load_workbook(stream)
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        self.assertIn("Roll Number", headers)
        self.assertIn("Subject Code", headers)
        self.assertIn("Status (PRESENT/ABSENT/EXCUSED)", headers)

    def test_parse_valid_excel_data(self):
        """Creates an in-memory workbook and validates parsing succeeds."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Roll Number", "Subject Code", "Date", "Period", "Status", "Remarks"])
        ws.append(["ST001", "AI101", "2026-10-21", 1, "PRESENT", "Good attendance"])

        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)

        result = parse_and_validate_attendance_excel(stream)
        self.assertTrue(result['success'])
        self.assertEqual(result['valid_count'], 1)
        self.assertEqual(result['error_count'], 0)

    def test_parse_detects_unknown_student_and_invalid_period(self):
        """Detects missing student roll number and out-of-range period."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Roll Number", "Subject Code", "Date", "Period", "Status", "Remarks"])
        ws.append(["NON_EXISTENT_ROLL", "AI101", "2026-10-21", 15, "PRESENT", ""])

        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)

        result = parse_and_validate_attendance_excel(stream)
        self.assertTrue(result['success'])
        self.assertEqual(result['valid_count'], 0)
        self.assertEqual(result['error_count'], 1)
        self.assertIn("not found in registered records", result['error_rows'][0]['errors'][0])
